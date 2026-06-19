"""Targeted fire-label reconciliation in labels.xlsx, in two directions:

  * SET   (fire 0 -> 1): scenes where the xlsx missed real flames YOLO boxed.
                         Scope: 3pplfire, onepersonfire.
  * CLEAR (fire 1 -> 0): scenes where the xlsx marked a hot-but-not-flame source
                         (a hairdryer) that YOLO correctly did NOT box.
                         Scope: 2pplwithhairdryer.

Both directions are gated on the YOLO ground truth for the exact (channel, frame).
Backs up labels.xlsx once (pristine pre-edit copy) and prints every cell changed.
Re-run scripts/check_label_agreement.py afterwards to confirm.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from thermal_algorithms.training.label_io import FIRE_CLASS_ID

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DATASET = Path("dataset")
XLSX = Path("dataset_raw/labels.xlsx")
BACKUP = Path("dataset_raw/labels.pre_fire_fix.xlsx")
SET_FIRE_SCENES = ["3pplfire", "onepersonfire"]      # fire 0 -> 1 where YOLO has fire
CLEAR_FIRE_SCENES = ["2pplwithhairdryer"]            # fire 1 -> 0 where YOLO has no fire


def yolo_has_fire(scene: str, ch: int, frame: int) -> bool:
    txt = DATASET / scene / f"ch{ch}_frames" / f"frame_{frame:05d}.txt"
    if not txt.is_file():
        return False
    for line in txt.read_text().splitlines():
        line = line.strip()
        if line and int(line.split()[0]) == FIRE_CLASS_ID:
            return True
    return False


def col_index(ws, name: str) -> int:
    """1-based column index of a header (row 1) by name."""
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if isinstance(cell.value, str) and cell.value.strip().lower() == name.lower():
            return cell.column
    raise KeyError(f"column {name!r} not found in sheet {ws.title!r}")


def main() -> None:
    if not XLSX.is_file():
        raise SystemExit(f"{XLSX} not found")
    if not BACKUP.exists():
        shutil.copy2(XLSX, BACKUP)
        print(f"Backup written: {BACKUP}")
    else:
        print(f"Backup already exists (kept): {BACKUP}")

    wb = openpyxl.load_workbook(XLSX)  # read-write, preserves all sheets/values
    total = 0

    def reconcile(scene: str, target_when_fire: bool) -> int:
        """If target_when_fire: set 0->1 where YOLO has fire.
        Else: set 1->0 where YOLO has NO fire. Returns #cells changed."""
        if scene not in wb.sheetnames:
            print(f"  WARN: sheet {scene!r} missing, skipping")
            return 0
        ws = wb[scene]
        c_ch = col_index(ws, "channel")
        c_fr = col_index(ws, "frame")
        c_fire = col_index(ws, "fire")
        changed = 0
        for row in ws.iter_rows(min_row=2):
            ch_cell, fr_cell, fire_cell = row[c_ch - 1], row[c_fr - 1], row[c_fire - 1]
            if ch_cell.value is None or fr_cell.value is None:
                continue
            try:
                ch = int(ch_cell.value); fr = int(fr_cell.value)
            except (ValueError, TypeError):
                continue
            try:
                cur = int(fire_cell.value) if fire_cell.value not in (None, "") else 0
            except (ValueError, TypeError):
                cur = 0
            has_fire = yolo_has_fire(scene, ch, fr)
            if target_when_fire and cur == 0 and has_fire:
                fire_cell.value = 1
                changed += 1
                print(f"  {scene} ch{ch} f{fr}: fire 0 -> 1")
            elif (not target_when_fire) and cur == 1 and not has_fire:
                fire_cell.value = 0
                changed += 1
                print(f"  {scene} ch{ch} f{fr}: fire 1 -> 0")
        verb = "set to fire=1" if target_when_fire else "cleared to fire=0"
        print(f"{scene}: {changed} cells {verb}")
        return changed

    for scene in SET_FIRE_SCENES:
        total += reconcile(scene, target_when_fire=True)
    for scene in CLEAR_FIRE_SCENES:
        total += reconcile(scene, target_when_fire=False)

    wb.save(XLSX)
    print(f"\nSaved {XLSX}. Total cells changed: {total}")


if __name__ == "__main__":
    main()
