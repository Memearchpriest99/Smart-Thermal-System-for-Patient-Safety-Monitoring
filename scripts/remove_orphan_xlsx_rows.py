"""Remove orphan xlsx rows — (channel, frame) rows whose thermal frame no longer
exists in dataset/ (frame index >= that scene's per-channel length).

Created to clean up after truncating per-channel frame-count mismatches. Scoped
by default to the two truncated scenes; pass scene names to override:
    python scripts/remove_orphan_xlsx_rows.py "personrunning"

Safety: aborts if Excel has the file open; snapshots labels.xlsx before editing;
only deletes rows whose frame index is out of range for its channel. Never edits
cell values, touch, CORRUPT, or tab colors.
"""

from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DATASET = Path("dataset")
XLSX = Path("dataset_raw/labels.xlsx")
LOCK = XLSX.with_name("~$" + XLSX.name)
DEFAULT_SCENES = ["personrunning", "setup2_1_man_running"]


def _resolve_scenes(argv: list[str]) -> list[str]:
    """--all [--skip A B ...] processes every scene (minus skips); otherwise the
    explicit scene names given; otherwise the default truncated pair."""
    if "--all" in argv:
        skip = set(argv[argv.index("--skip") + 1:]) if "--skip" in argv else set()
        return sorted(
            d.name for d in DATASET.iterdir()
            if d.is_dir() and (d / "ch0_raw_data.npz").is_file() and d.name not in skip
        )
    return argv or DEFAULT_SCENES


SCENES = _resolve_scenes(sys.argv[1:])


def main() -> None:
    if LOCK.exists():
        raise SystemExit("ABORT: Excel has labels.xlsx open (~$ lock present). Close it first.")

    # Per-scene per-channel valid length from meta.json.
    lengths = {}
    for scene in SCENES:
        n = json.loads((DATASET / scene / "meta.json").read_text())["n_frames_per_ch"]
        lengths[scene] = {int(k): int(v) for k, v in n.items()}

    from snapshot_util import snapshot_xlsx
    snap = snapshot_xlsx(XLSX)
    print(f"Snapshot: {snap}")

    wb = openpyxl.load_workbook(XLSX)
    total = 0
    for scene in SCENES:
        if scene not in wb.sheetnames:
            print(f"  WARN: sheet {scene!r} missing"); continue
        ws = wb[scene]
        hdr = {c.value: c.column for c in next(ws.iter_rows(min_row=1, max_row=1))}
        c_ch, c_fr = hdr["channel"], hdr["frame"]
        # Collect 1-based worksheet row numbers to delete.
        to_delete = []
        for row in ws.iter_rows(min_row=2):
            ch_v = row[c_ch - 1].value
            fr_v = row[c_fr - 1].value
            if ch_v is None or fr_v is None:
                continue
            try:
                ch = int(ch_v); fr = int(fr_v)
            except (ValueError, TypeError):
                continue
            if fr >= lengths[scene].get(ch, 10 ** 9):
                to_delete.append((row[0].row, ch, fr))
        # Delete bottom-up so row numbers stay valid.
        for rownum, ch, fr in sorted(to_delete, reverse=True):
            ws.delete_rows(rownum, 1)
            print(f"  {scene}: removed row (ch{ch} f{fr})")
        print(f"{scene}: {len(to_delete)} orphan rows removed")
        total += len(to_delete)

    wb.save(XLSX)
    print(f"\nSaved {XLSX}. Total orphan rows removed: {total}")


if __name__ == "__main__":
    main()
