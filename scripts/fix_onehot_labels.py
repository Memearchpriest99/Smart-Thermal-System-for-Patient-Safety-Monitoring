"""Fix the 2 xlsx one-hot violations (a frame had >1 human-count column set).

Each target was disambiguated against YOLO box counts AND the surrounding xlsx
rows, then the stray bit is cleared so exactly one of single/two/three remains:

  * 3ppl ch0 f47        : two=1 AND three=1  -> clear three (run of "two"; YOLO=2)
  * onepersonfire ch2 f75: single=1 AND three=1 -> clear three (one-person scene; YOLO=1)

Guarded: each cell must currently equal 1 before it is cleared, else the script
aborts (protects against re-running on already-changed / unexpected data).
The pristine original workbook is preserved at labels.pre_fire_fix.xlsx.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

XLSX = Path("dataset_raw/labels.xlsx")

# (scene, channel, frame, column-header-to-clear)
TARGETS = [
    ("3ppl", 0, 47, "three humans"),
    ("onepersonfire", 2, 75, "three humans"),
]


def header_map(ws) -> dict:
    return {c.value: c.column for c in next(ws.iter_rows(min_row=1, max_row=1))}


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    changed = 0
    for scene, ch, fr, col_name in TARGETS:
        ws = wb[scene]
        h = header_map(ws)
        c_ch, c_fr, c_col = h["channel"], h["frame"], h[col_name]
        hit = False
        for row in ws.iter_rows(min_row=2):
            if (row[c_ch - 1].value is not None and int(row[c_ch - 1].value) == ch
                    and row[c_fr - 1].value is not None and int(row[c_fr - 1].value) == fr):
                cell = row[c_col - 1]
                cur = int(cell.value) if cell.value not in (None, "") else 0
                if cur != 1:
                    raise SystemExit(
                        f"ABORT: {scene} ch{ch} f{fr} {col_name!r} = {cur} (expected 1). "
                        f"Data already changed or unexpected — no edits saved."
                    )
                cell.value = 0
                changed += 1
                hit = True
                print(f"  {scene} ch{ch} f{fr}: {col_name!r} 1 -> 0")
                break
        if not hit:
            raise SystemExit(f"ABORT: row not found for {scene} ch{ch} f{fr}")

    wb.save(XLSX)
    print(f"\nSaved {XLSX}. Cells cleared: {changed}")


if __name__ == "__main__":
    main()
