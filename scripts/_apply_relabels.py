"""Apply the consensus + human-reviewed touch labels to the 3 re-annotated videos.

Merges reports/auto_label_proposals.json (agreements) with Guy's 167 review
decisions, then writes:
  - datasets/waveshare_work/labels.xlsx   (touch column, all 3 channel rows/frame)
  - datasets/waveshare_work/contact_labels.csv
Backups were taken beforehand in datasets/_label_backups/.
Tab colors are left untouched (machine-assisted labels, not human-verified).
"""
import csv, json, sys
from pathlib import Path
_ROOT = Path(__file__).resolve().parent.parent
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
import openpyxl

DECISIONS = Path(sys.argv[1])
XLSX = _ROOT / "datasets" / "waveshare_work" / "labels.xlsx"
CSV = _ROOT / "datasets" / "waveshare_work" / "contact_labels.csv"

props = json.loads((_ROOT / "reports" / "auto_label_proposals.json").read_text())
dec = {}
for row in csv.DictReader(open(DECISIONS)):
    dec[(row["scene"], int(row["frame_idx"]))] = int(row["contact"])

final = {}   # scene -> {fi: label}
for scene, recs in props.items():
    final[scene] = {}
    for r in recs:
        if r["agree"]:
            final[scene][r["fi"]] = int(r["proposed"])
        else:
            final[scene][r["fi"]] = dec[(scene, r["fi"])]

# ---- xlsx: touch column on every channel row of each frame -------------------
wb = openpyxl.load_workbook(XLSX)
for scene, labels in final.items():
    ws = wb[scene]
    hdr = [c.value for c in ws[1]]
    fcol = hdr.index("frame") + 1
    tcol = hdr.index("touch") + 1
    changed = 0
    for row in ws.iter_rows(min_row=2):
        fi = row[fcol - 1].value
        if fi is None: continue
        fi = int(fi)
        if fi in labels and row[tcol - 1].value != labels[fi]:
            row[tcol - 1].value = labels[fi]
            changed += 1
    print(f"{scene}: {changed} xlsx cells updated")
wb.save(XLSX)
print(f"saved {XLSX}")

# ---- contact_labels.csv: replace values for the 3 scenes ----------------------
rows = list(csv.DictReader(open(CSV)))
n_upd = 0
for row in rows:
    sc = row["session"]
    if sc in final and int(row["frame_idx"]) in final[sc]:
        new = str(final[sc][int(row["frame_idx"])])
        if row["contact"] != new:
            row["contact"] = new
            n_upd += 1
with open(CSV, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["session", "frame_idx", "contact"])
    w.writeheader(); w.writerows(rows)
print(f"contact_labels.csv: {n_upd} rows updated")

for scene, labels in final.items():
    pos = sum(labels.values())
    print(f"  {scene}: {pos} positive / {len(labels)} frames")
