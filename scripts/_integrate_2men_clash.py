"""Integrate the new '2men_clash' session (Guy's freshly-labelled touch data)
into datasets/waveshare_work.

Source: Desktop/Final Project/More data/DAR&BASH/DAR&BASH/2men_clash/20260630_162507
 - ch0_thermal.npz has 367 frames; PNG/label correlation proves the labels map
   to frames[0:355] -> trim the 12 trailing frames. ch1/ch2 are already 355.
 - Renames ch{k}_thermal.npz -> ch{k}_raw_data.npz (key 'frames'), copies
   ch{k}_frames/, normalizes classes.txt to "fire\nperson", writes meta.json.
 - Appends 355 rows to the consolidated contact_labels.csv (74 positives).
 - Adds a '2men_clash' sheet to labels.xlsx (root workbook snapshotted first).
"""
import csv, json, shutil, sys
from pathlib import Path
_ROOT = Path(__file__).resolve().parent.parent
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
import numpy as np
import openpyxl

SRC = Path(r"C:\Users\Guy\Desktop\Final Project\More data\DAR&BASH\DAR&BASH\2men_clash\20260630_162507")
DST = _ROOT / "datasets" / "waveshare_work" / "2men_clash"
N = 355

assert not DST.exists(), f"{DST} already exists — refusing to overwrite"
DST.mkdir(parents=True)

# frames npz (trim ch0), pngs/txts
for ch in range(3):
    arr = np.load(SRC / f"ch{ch}_thermal.npz")["frames"][:N].astype(np.float32)
    assert arr.shape == (N, 62, 80), arr.shape
    np.savez_compressed(DST / f"ch{ch}_raw_data.npz", frames=arr)
    shutil.copytree(SRC / f"ch{ch}_frames", DST / f"ch{ch}_frames")
    print(f"ch{ch}: npz {arr.shape} + {len(list((DST / f'ch{ch}_frames').iterdir()))} frame files")

(DST / "classes.txt").write_text("fire\nperson\n")
(DST / "meta.json").write_text(json.dumps({
    "scene": "2men_clash", "sensor": "waveshare", "setup": None,
    "original_timestamp": "20260630_162507", "fps": 8.0,
    "n_frames_per_ch": {"0": N, "1": N, "2": N},
    "duplicate_of": None,
    "alignment_fixes": ["ch0_thermal.npz had 367 frames; trimmed to [0:355] "
                        "(PNG correlation: labels correspond to frames[0:355])"],
    "truncations": [],
}, indent=2))

# labels: per-session csv -> consolidated csv
src_rows = list(csv.DictReader(open(SRC / "contact_labels.csv")))
assert len(src_rows) == N, len(src_rows)
cons = _ROOT / "datasets" / "waveshare_work" / "contact_labels.csv"
rows = list(csv.DictReader(open(cons)))
assert not any(r["session"] == "2men_clash" for r in rows)
for r in src_rows:
    rows.append({"session": "2men_clash", "frame_idx": r["frame_idx"], "contact": r["contact"]})
with open(cons, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["session", "frame_idx", "contact"])
    w.writeheader(); w.writerows(rows)
pos = sum(int(r["contact"]) for r in src_rows)
print(f"contact_labels.csv: +{len(src_rows)} rows ({pos} positive)")

# labels.xlsx: snapshot, then copy the session's sheet into the root workbook
root_x = _ROOT / "datasets" / "waveshare_work" / "labels.xlsx"
bkp = _ROOT / "datasets" / "_label_backups" / "labels_pre_2men_clash_2026-07-02.xlsx"
shutil.copy2(root_x, bkp)
src_wb = openpyxl.load_workbook(SRC / "labels.xlsx")
dst_wb = openpyxl.load_workbook(root_x)
src_ws = src_wb[src_wb.sheetnames[0]]
assert "2men_clash" not in dst_wb.sheetnames
dst_ws = dst_wb.create_sheet("2men_clash")
for row in src_ws.iter_rows(values_only=True):
    dst_ws.append(row)
dst_wb.save(root_x)
print(f"labels.xlsx: added sheet '2men_clash' ({src_ws.max_row - 1} rows); backup at {bkp.name}")
